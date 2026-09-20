"""CNT3054 청약 인식 프로젝트(B안): 분류 신뢰도 검증 (Human-in-the-loop)

LLM(gpt-oss-120b) 분류를 사람이 직접 한 분류와 비교해 얼마나 믿을 만한지 확인한다.
  - Cohen's κ(카파): 우연히 맞을 확률을 뺀 일치도. 해석 기준 (Landis & Koch, 1977)
      0.41~0.60 보통 / 0.61~0.80 상당한 일치 / 0.81 이상 거의 완전한 일치
    https://doi.org/10.2307/2529310
  - scikit-learn cohen_kappa_score:
    https://scikit-learn.org/stable/modules/generated/sklearn.metrics.cohen_kappa_score.html

진행 순서
  (1) 코딩 시트 만들기 — 분석 기간(2026-09-13~19) 댓글 중 무작위 100개. LLM 결과는 넣지 않는다.
      python check_reliability.py sample
      → output/reliability_sheet.xlsx
        · "코딩" 시트: 태도·집값 전망·통장 행동은 드롭다운으로 고르고, 관심사는 A~I를 붙여 쓴다(예: AB)
        · "기준" 시트: cy-index.txt 분류 기준 전문
  (2) 조원 2명이 파일을 복사해 각자 따로 채운다 (서로의 답, LLM 결과를 보지 않기)
      예: output/reliability_sheet_kim.xlsx, output/reliability_sheet_lee.xlsx
  (3) 일치도 계산 — 사람 vs LLM, 사람 vs 사람을 함께 계산하고 불일치 목록을 저장
      python check_reliability.py score output/reliability_sheet_kim.xlsx output/reliability_sheet_lee.xlsx
      → 화면에 κ 표, output/reliability_disagreements.csv 에 LLM과 다르게 판단한 댓글 목록
  (4) 불일치 댓글을 조원이 함께 보고 합의한 정답을 정하면, 보고서에 κ와 함께 대표적인 오분류 유형을 적는다.
      오분류가 한쪽으로 몰리면(예: 양가를 부정으로) cy-index.txt 기준을 고치고 classify_comments.py 를 다시 돌린다.
"""
import argparse
from itertools import combinations
from pathlib import Path

import pandas as pd
from sklearn.metrics import cohen_kappa_score

BASE = Path(__file__).resolve().parent
TOPIC_CODES = "ABCDEFGHI"   # cy-index.txt 의 관심사 코드
COLS = ["stance", "topics", "outlook", "action"]
KOR = {"stance": "태도", "topics": "관심사", "outlook": "집값 전망", "action": "통장 행동"}
PERIOD_START, PERIOD_END = "2026-09-13", "2026-09-19"   # 분석 기간 (한국 시간)
CHOICES = {"stance": ["N", "M", "P", "X"], "outlook": ["0", "U", "D"], "action": ["0", "H", "K"]}


def latest_classified():
    runs = sorted((BASE / "output").glob("*/classified.csv"), key=lambda p: p.stat().st_mtime)
    if not runs:
        raise SystemExit("output/ 에 classified.csv 가 없습니다. classify_comments.py 를 먼저 실행하세요.")
    return runs[-1]


def read_sheet(path):
    """사람이 채운 코딩 시트(.xlsx 또는 .csv)를 읽어 comment_id 기준 표로 만든다."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        df = pd.read_excel(path, sheet_name="코딩", dtype=str)
    else:
        df = pd.read_csv(path, dtype=str)
    df = df.rename(columns={"태도": "stance", "관심사": "topics", "집값 전망": "outlook", "통장 행동": "action"})
    df = df.dropna(subset=["stance"]).set_index("comment_id")
    df["topics"] = df["topics"].fillna("").str.upper().str.replace(r"[^A-I]", "", regex=True)
    for c in ("outlook", "action"):
        df[c] = df[c].fillna("0").str.strip()
    df["stance"] = df["stance"].str.strip().str.upper()
    return df[COLS]


def read_llm(path):
    df = pd.read_csv(path, dtype=str).set_index("comment_id")
    df["topics"] = df["topics"].fillna("")
    for c in ("outlook", "action"):
        df[c] = df[c].fillna("0")
    return df


def sample(src, n, seed):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    d = pd.read_csv(src)
    day = pd.to_datetime(d["published_at"], utc=True).dt.tz_convert("Asia/Seoul").dt.strftime("%Y-%m-%d")
    d = d[(day >= PERIOD_START) & (day <= PERIOD_END)]
    s = d.sample(n=min(n, len(d)), random_state=seed)

    wb = Workbook()
    ws = wb.active
    ws.title = "코딩"
    head = ["번호", "comment_id", "댓글", "태도", "관심사", "집값 전망", "통장 행동", "메모"]
    ws.append(head)
    for i, r in enumerate(s.itertuples(), 1):
        ws.append([i, r.comment_id, str(r.text_raw), None, None, None, None, None])
    widths = [6, 14, 70, 8, 10, 10, 10, 24]
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w
    fill = PatternFill("solid", fgColor="E8EEF7")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), fill
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.column_dimensions["B"].hidden = True   # 연결용 ID 는 숨김 (지우지 말 것)
    last = ws.max_row
    for col, key, prompt in [("D", "stance", "N 부정 / M 양가 / P 긍정 / X 해당 없음"),
                             ("F", "outlook", "0 없음 / U 상승 / D 하락"),
                             ("G", "action", "0 없음 / H 해지 / K 유지")]:
        dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES[key]) + '"', allow_blank=False,
                            promptTitle=KOR[key], prompt=prompt, showInputMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    tv = DataValidation(type="custom", formula1='=OR(E2="",ISNUMBER(SEARCH(LEFT(E2,1),"ABCDEFGHI")))',
                        promptTitle="관심사", showInputMessage=True,
                        prompt="해당하는 코드를 모두 붙여 쓰기 (예: AB). 없으면 비워 두기")
    ws.add_data_validation(tv)
    tv.add(f"E2:E{last}")

    guide = wb.create_sheet("기준")
    guide.column_dimensions["A"].width = 110
    guide.append(["작성 방법: 다른 조원의 답이나 LLM 결과를 보지 말고 혼자 판단하세요. 파일 이름에 본인 이름을 붙여 저장하세요."])
    guide.append([""])
    for line in (BASE / "cy-index.txt").read_text(encoding="utf-8").splitlines():
        if not line.startswith("#"):
            guide.append([line])

    out = BASE / "output" / "reliability_sheet.xlsx"
    wb.save(out)
    print(f"저장: {out} ({len(s)}개, seed={seed}, 기간 {PERIOD_START}~{PERIOD_END})")
    print("조원마다 파일을 복사해 이름을 바꾸고(예: reliability_sheet_kim.xlsx) 각자 채운 뒤 score 명령으로 비교하세요.")


def kappa_row(a, b):
    """두 분류 결과(같은 comment_id)의 항목별 κ 와 단순 일치율."""
    j = a.join(b, rsuffix="_2", how="inner")
    row = {"n": len(j)}
    for c in ["stance", "outlook", "action"]:
        x, y = j[c], j[c + "_2"]
        row[f"{KOR[c]} κ"] = cohen_kappa_score(x, y) if len(set(x) | set(y)) > 1 else 1.0
        row[f"{KOR[c]} 일치"] = (x == y).mean()
    ks = []
    for code in TOPIC_CODES:   # 관심사는 복수 선택 → 코드마다 있다/없다를 비교해 평균
        x, y = j["topics"].str.contains(code), j["topics_2"].str.contains(code)
        if x.any() or y.any():
            ks.append(cohen_kappa_score(x, y))
    row["관심사 κ(평균)"] = sum(ks) / len(ks) if ks else float("nan")
    return row


def score(files, ref):
    llm = read_llm(ref)
    people = {Path(f).stem.replace("reliability_sheet_", ""): read_sheet(f) for f in files}
    rows = {}
    for name, h in people.items():
        rows[f"{name} vs LLM"] = kappa_row(h, llm[COLS])
    for (n1, h1), (n2, h2) in combinations(people.items(), 2):
        rows[f"{n1} vs {n2}"] = kappa_row(h1, h2)
    table = pd.DataFrame(rows).T
    pd.set_option("display.width", 200)
    print(f"기준(LLM): {ref}")
    print(table.round(3).to_string())
    print("\nκ 해석: 0.41~0.60 보통 / 0.61~0.80 상당한 일치 / 0.81 이상 거의 완전한 일치")

    # LLM 과 다르게 판단한 댓글 목록 (조원 합의용)
    text = pd.read_csv(ref, dtype=str).set_index("comment_id")["text_raw"]
    recs = []
    for name, h in people.items():
        j = h.join(llm[COLS], rsuffix="_llm", how="inner")
        for cid, r in j.iterrows():
            diff = [KOR[c] for c in COLS if r[c] != r[c + "_llm"]]
            if diff:
                recs.append({"comment_id": cid, "코더": name, "다른 항목": ", ".join(diff), "댓글": text.get(cid, ""),
                             **{f"{KOR[c]}(사람)": r[c] for c in COLS}, **{f"{KOR[c]}(LLM)": r[c + "_llm"] for c in COLS}})
    out = BASE / "output" / "reliability_disagreements.csv"
    pd.DataFrame(recs).sort_values(["comment_id", "코더"]).to_csv(out, index=False, encoding="utf-8-sig")
    table.round(3).to_csv(BASE / "output" / "reliability_kappa.csv", encoding="utf-8-sig")
    print(f"\n저장: {out} ({len(recs)}건), output/reliability_kappa.csv")


def main():
    ap = argparse.ArgumentParser(description="분류 신뢰도 검증")
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("sample")
    a.add_argument("--input", help="classified.csv (기본: output/ 의 가장 최근 파일)")
    a.add_argument("--n", type=int, default=100)
    a.add_argument("--seed", type=int, default=42)
    b = sub.add_parser("score")
    b.add_argument("files", nargs="+", help="사람이 채운 코딩 시트 (.xlsx 또는 .csv)")
    b.add_argument("--ref", help="LLM 분류 결과 (기본: output/ 의 가장 최근 classified.csv)")
    args = ap.parse_args()
    if args.cmd == "sample":
        sample(args.input or latest_classified(), args.n, args.seed)
    else:
        score(args.files, args.ref or latest_classified())


if __name__ == "__main__":
    main()
